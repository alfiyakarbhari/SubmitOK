import os
import re
from docx import Document
import pdfplumber

# =====================================================
# SECTION KEYWORDS (GENERALIZED)
# =====================================================

SECTION_KEYWORDS = {
    "aim": [
        "aim", "objective", "objectives", "problem statement"
    ],
    "theory": [
        "theory"
    ],
    "code": [
        "code", "program", "server", "client"
    ],
    "output": [
        "output", "result"
    ],
    "methodology": [
        "methodology", "method", "procedure",
        "working", "steps", "algorithm"
    ],
    "conclusion": [
        "conclusion"
    ],
    "student_details": [
        "name", "uid", "usn", "roll",
        "experiment", "exp"
    ]
}

# =====================================================
# UTILS
# =====================================================

def normalize(text):
    text = text.lower()
    text = re.sub(r"[^a-z0-9 ]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def detect_section(line):
    normalized = normalize(line)
    for section, keywords in SECTION_KEYWORDS.items():
        for kw in keywords:
            if normalized.startswith(kw):
                return section
    return None


def word_count(text):
    return len(text.split())


# =====================================================
# TEXT EXTRACTION
# =====================================================

def extract_text_from_docx(file_path):
    doc = Document(file_path)
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def extract_text_from_pdf(file_path):
    lines = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                for line in text.split("\n"):
                    line = line.strip()
                    if line:
                        lines.append(line)
    return lines


# =====================================================
# IMAGE DETECTION
# =====================================================

def docx_has_images(file_path):
    doc = Document(file_path)
    return len(doc.inline_shapes) > 0


def pdf_has_images(file_path):
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            if page.images:
                return True
    return False


# =====================================================
# PARSING DOCUMENT (FIXED)
# =====================================================

def parse_document(file_path):
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".docx":
        lines = extract_text_from_docx(file_path)
        has_image = docx_has_images(file_path)
    elif ext == ".pdf":
        lines = extract_text_from_pdf(file_path)
        has_image = pdf_has_images(file_path)
    else:
        raise ValueError("Unsupported file type")

    sections = {}
    current_section = None
    parsing_started = False   # 🔑 prevents student_details hijack

    for line in lines:
        detected = detect_section(line)

        # Ignore student details after real sections begin
        if detected == "student_details" and parsing_started:
            detected = None

        if detected and detected != "student_details":
            current_section = detected
            sections[current_section] = ""
            parsing_started = True
        elif current_section:
            sections[current_section] += line + " "

    full_text = " ".join(lines)
    return sections, full_text, has_image


# =====================================================
# STUDENT DOCUMENT JSON
# =====================================================

def generate_student_document_json(sections, full_text, has_image):
    return {
        "aim": {
            "present": "aim" in sections,
            "word_count": word_count(sections["aim"]) if "aim" in sections else 0
        },
        "theory": {
            "present": "theory" in sections,
            "word_count": word_count(sections["theory"]) if "theory" in sections else 0
        },
        "code": {
            "present": "code" in sections,
            "comment_lines": 0
        },
        "output": {
            "present": "output" in sections,
            "has_image": has_image
        },
        "methodology": {
            "present": "methodology" in sections
        },
        "conclusion": {
            "present": "conclusion" in sections,
            "word_count": word_count(sections["conclusion"]) if "conclusion" in sections else 0
        },
        "full_text": full_text
    }


# =====================================================
# RUN
# =====================================================

if __name__ == "__main__":
    file_path = r"C:\Users\SUVARNA HOUSE\Downloads\DC\DC_EXP_3.pdf"

    sections, full_text, has_image = parse_document(file_path)

    student_document_json = generate_student_document_json(
        sections, full_text, has_image
    )

    print(student_document_json)
